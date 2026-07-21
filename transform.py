import json
import os
import re
import sys
from datetime import datetime
from pathlib import Path

import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE_DIR = Path(__file__).parent

COLORS = {
    'header_bg':   '1F4E79',
    'header_font': 'FFFFFF',
    'city_font':   'A0B4C8',
    'phone_font':  'D4E4F7',
    'tg_font':     '64B5F6',
    'group_bg':    'BDD7EE',
    'group_font':  '1F4E79',
    'alt_row':     'EBF3FB',
}

COL_WIDTHS = {'A': 7.57, 'B': 13.14, 'C': 14.43, 'D': 77.29, 'E': 16.0, 'F': 15.57}


def load_config():
    with open(BASE_DIR / 'config.json', encoding='utf-8') as f:
        return json.load(f)


def _side():
    return Side(style='thin', color='B8CCE4')


def _border():
    s = _side()
    return Border(left=s, right=s, top=s, bottom=s)


def _apply(cell, style):
    for attr, val in style.items():
        setattr(cell, attr, val)


# ─── Чтение порядка групп из файла ────────────────────────────────────────────

def load_group_order(config):
    order_file = BASE_DIR / config.get('group_order_file', 'Места_по_группам_товара_исправлено.xlsx')
    if not order_file.exists():
        return {}

    df = pd.read_excel(order_file, skiprows=1, header=None)
    order = {}
    for _, row in df.iterrows():
        group_name = str(row.iloc[1]).strip() if pd.notna(row.iloc[1]) else ''
        position = row.iloc[2] if len(row) > 2 and pd.notna(row.iloc[2]) else None
        if group_name and group_name != 'Группа товара' and position is not None:
            try:
                order[group_name] = int(float(position))
            except (ValueError, TypeError):
                pass
    return order


# ─── Чтение наценок из файла ──────────────────────────────────────────────────

def load_markups(config):
    markups_path = BASE_DIR / config.get('markups_file', 'наценки по группам товара.xlsx')
    if not markups_path.exists():
        return {}, {}

    df = pd.read_excel(markups_path, skiprows=2, header=None)

    group_markups = {}
    homeline_prices = {}
    in_homeline = False

    for _, row in df.iterrows():
        col_b = str(row.iloc[1]).strip() if pd.notna(row.iloc[1]) else ''
        col_c = row.iloc[2] if len(row) > 2 else None

        if col_b == 'Артикул':
            in_homeline = True
            continue

        if not in_homeline:
            if col_b and col_b not in ('Группа товара', 'Наценка (%)', 'nan'):
                try:
                    markup_val = int(float(col_c)) if pd.notna(col_c) else 0
                    group_markups[col_b] = markup_val
                except (ValueError, TypeError):
                    pass
        else:
            article = col_b
            col_d = row.iloc[3] if len(row) > 3 else None
            if article and article != 'nan' and col_d is not None and pd.notna(col_d):
                try:
                    homeline_prices[article] = float(col_d)
                except (ValueError, TypeError):
                    pass

    return group_markups, homeline_prices


# ─── Сортировка генераторов ───────────────────────────────────────────────────

def _generator_sort_key(name: str) -> tuple:
    """Ключ сортировки внутри группы генераторов.
    Порядок: инверт.бенз → инверт.диз → инверт.двухтопл → бенз → диз → двухтопл → автоматика/прочее
    Внутри каждой подгруппы — по мощности по возрастанию.
    """
    n = name.lower()
    is_inverter = 'инвертор' in n
    is_diesel   = 'дизел' in n
    is_dual     = 'duomatic' in n or 'двухтопл' in n or 'дуомат' in n
    is_avr      = 'авр' in n or ' ats' in n or 'автоматик' in n

    if is_avr:
        subtype = 6
    elif is_inverter and is_dual:
        subtype = 2
    elif is_inverter and is_diesel:
        subtype = 1
    elif is_inverter:
        subtype = 0
    elif is_dual:
        subtype = 5
    elif is_diesel:
        subtype = 4
    else:
        subtype = 3  # бензиновый не инверторный

    numbers = [int(m) for m in re.findall(r'\d+', name) if 500 <= int(m) <= 25000]
    power = numbers[0] if numbers else 99999

    return (subtype, power)


# ─── Извлечение бренда из названия товара ────────────────────────────────────

def extract_brand(name: str) -> str:
    """Возвращает бренд из названия товара.
    Бренд — первое слово из заглавных букв (≥4 символов), не являющееся
    русским прилагательным (окончания ОЙ, АЯ, ОЕ, ЫЙ, ИЙ).
    Пример: 'Аккумуляторная дрель EINHELL PXC...' → 'EINHELL'
    Если не найден — возвращает 'Noname'.
    """
    _adj_endings = re.compile(r'(ОЙ|АЯ|ОЕ|ЫЙ|ИЙ|УЮ|ЫЕ|ИЕ)$')
    for word in name.split():
        letters = re.sub(r'[^а-яёА-ЯЁa-zA-Z]', '', word)
        if len(letters) >= 4 and letters == letters.upper():
            if _adj_endings.search(letters):
                continue
            return word.rstrip(',').rstrip('.')
    return 'Noname'


# ─── Чтение прайса поставщика ─────────────────────────────────────────────────
#
# Прайс ИП Аксёнов приходит в двух макетах (структура «плавает» от версии к версии):
#   • табличный — 1 строка = 1 товар (колонки наименования/цены ищутся по шапке);
#   • карточный — 1 товар = блок из нескольких строк, цена текстом «520,00 RUB».
# read_supplier_price сам определяет макет и вызывает нужный разборщик.
# Оба возвращают одинаковый DataFrame: [article, group, brand, name, price_in].
# Артикулы всегда генерируются (UT-000001…) — коды поставщика не выводятся.

# Ячейка-цена карточного формата: «1 140,00 RUB», «520,00 RUB» и т.п.
_PRICE_RUB_RE = r'\d[\d\s ]*[.,]\d{2}\s*RUB'
_NAME_KW = ('номенклатура', 'наименование', 'название')
_PRICE_KW = ('цена', 'опт', 'ррц', 'стоимост', 'прайс')


def _parse_price_text(value):
    """'1 140,00 RUB' → 1140.0. Возвращает None, если распознать не удалось."""
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return None
    s = re.sub(r'[^\d.,]', '', str(value))   # убираем пробелы, 'RUB', валюту
    if not s:
        return None
    s = s.replace(',', '.')                   # запятая — десятичный разделитель
    if s.count('.') > 1:                      # на случай нескольких точек
        head, _, tail = s.rpartition('.')
        s = head.replace('.', '') + '.' + tail
    try:
        return float(s)
    except ValueError:
        return None


def detect_layout(df):
    """Определяет макет по содержимому: 'card', если встречаются ячейки-цены
    вида «… RUB» (карточный), иначе 'table' (табличный). В табличном формате
    цены числовые, поэтому таких ячеек там нет вовсе."""
    for c in range(df.shape[1]):
        col = df.iloc[:, c].dropna().astype(str)
        if len(col) and col.str.contains(_PRICE_RUB_RE, regex=True, na=False).any():
            return 'card'
    return 'table'


def _find_header_row(df):
    """Ищет строку-заголовок (по ключевому слову наименования) в первых 25 строках.
    Возвращает (header_row, name_col, price_col); любой из них может быть None."""
    for r in range(min(25, len(df))):
        name_col = price_col = None
        for c in range(df.shape[1]):
            v = df.iat[r, c]
            if pd.isna(v):
                continue
            s = str(v).strip().lower()
            if name_col is None and any(k in s for k in _NAME_KW):
                name_col = c
            if price_col is None and any(k in s for k in _PRICE_KW):
                price_col = c
        if name_col is not None:
            return r, name_col, price_col
    return None, None, None


def _detect_price_col(df, start):
    """Колонка с наибольшим числом числовых значений в теле (fallback)."""
    best_col, best_cnt = None, 0
    for c in range(df.shape[1]):
        cnt = int(pd.to_numeric(df.iloc[start:, c], errors='coerce').notna().sum())
        if cnt > best_cnt:
            best_col, best_cnt = c, cnt
    return best_col if best_cnt > 0 else None


def _detect_name_col(df, start, exclude):
    """Колонка с наибольшим числом текстовых значений (fallback)."""
    best_col, best_cnt = None, 0
    for c in range(df.shape[1]):
        if c == exclude:
            continue
        col = df.iloc[start:, c]
        texts = int(col.notna().sum() - pd.to_numeric(col, errors='coerce').notna().sum())
        if texts > best_cnt:
            best_col, best_cnt = c, texts
    return best_col if best_cnt > 0 else None


def parse_table(df, config):
    """Табличный макет: 1 строка = 1 товар. Колонки наименования и цены
    определяются по строке-заголовку; строки без цены — заголовки групп."""
    header_row, name_col, price_col = _find_header_row(df)
    if header_row is not None:
        body_start = header_row + 1
        if price_col is None:
            price_col = _detect_price_col(df, body_start)
    else:
        body_start = config.get('read_options', {}).get('skiprows', 8)
        price_col = _detect_price_col(df, body_start)
        name_col = _detect_name_col(df, body_start, exclude=price_col)

    if name_col is None or price_col is None:
        raise ValueError(
            "Не удалось определить структуру табличного файла: "
            "не найдена колонка наименования или цены."
        )

    rows = []
    current_group = ''
    counter = 1
    for r in range(body_start, len(df)):
        nv = df.iat[r, name_col]
        name = str(nv).strip() if pd.notna(nv) else ''
        if not name or name == 'nan':
            continue
        pv = df.iat[r, price_col] if price_col < df.shape[1] else None
        price_num = pd.to_numeric(pv, errors='coerce')
        if pd.isna(price_num):
            current_group = name
        else:
            rows.append({
                'article':  f'UT-{counter:06d}',
                'group':    current_group,
                'brand':    extract_brand(name),
                'name':     name,
                'price_in': float(price_num),
            })
            counter += 1
    return pd.DataFrame(rows)


def _find_rub_col(df):
    """Колонка с наибольшим числом ячеек-цен «… RUB»."""
    best_col, best_cnt = None, 0
    for c in range(df.shape[1]):
        col = df.iloc[:, c].dropna().astype(str)
        cnt = int(col.str.contains(_PRICE_RUB_RE, regex=True, na=False).sum()) if len(col) else 0
        if cnt > best_cnt:
            best_col, best_cnt = c, cnt
    return best_col


def _find_label_col(df, label):
    """Колонка, в которой встречается ячейка-подпись (напр. 'Код')."""
    for c in range(df.shape[1]):
        col = df.iloc[:, c].dropna().astype(str).str.strip()
        if len(col) and (col == label).any():
            return c
    return None


def _find_col_containing(df, substr):
    """Колонка, где встречается ячейка с подстрокой (напр. 'ФОТОГРАФ')."""
    up = substr.upper()
    for c in range(df.shape[1]):
        col = df.iloc[:, c].dropna().astype(str).str.upper()
        if len(col) and col.str.contains(up, regex=False, na=False).any():
            return c
    return None


def parse_cards(df, config):
    """Карточный макет: 1 товар = блок строк. Якорь — строка со значением цены
    «… RUB». Наименование берётся из строки с фото-плейсхолдером выше, заголовки
    групп — из колонки фото (текст, не «НЕТ ФОТОГРАФИИ»). Подписи Код/Артикул/Цена
    пропускаются."""
    price_col = _find_rub_col(df)
    if price_col is None:
        raise ValueError("Карточный формат: не найдена колонка цены «… RUB».")

    # Колонка наименования = где стоит подпись 'Код' (там же наименования и коды);
    # fallback — самая текстовая колонка.
    name_col = _find_label_col(df, 'Код')
    if name_col is None:
        name_col = _detect_name_col(df, 0, exclude=price_col)
    group_col = _find_col_containing(df, 'ФОТОГРАФ')

    _labels = ('Код', 'Артикул', 'Цена')
    rows = []
    current_group = ''
    pending_name = None
    counter = 1

    for r in range(len(df)):
        pv = df.iat[r, price_col] if price_col < df.shape[1] else None
        pv_s = str(pv) if pd.notna(pv) else ''
        price = _parse_price_text(pv) if 'RUB' in pv_s.upper() else None

        if price is not None:
            rows.append({
                'article':  f'UT-{counter:06d}',
                'group':    current_group,
                'brand':    extract_brand(pending_name or ''),
                'name':     pending_name or '',
                'price_in': price,
            })
            counter += 1
            pending_name = None
            continue

        gv = df.iat[r, group_col] if (group_col is not None and group_col < df.shape[1]) else None
        group_val = str(gv).strip() if pd.notna(gv) else ''
        if group_val and 'ФОТОГРАФ' not in group_val.upper():
            current_group = group_val
            pending_name = None
            continue

        nv = df.iat[r, name_col] if name_col is not None and name_col < df.shape[1] else None
        name_val = str(nv).strip() if pd.notna(nv) else ''
        if name_val and name_val not in _labels and 'ФОТОГРАФ' not in name_val.upper():
            if pending_name is None:      # первая строка блока = наименование
                pending_name = name_val

    return pd.DataFrame(rows)


def read_supplier_price(filepath, config):
    """Точка входа: определяет макет файла и разбирает его соответствующим
    разборщиком. Возвращает DataFrame [article, group, brand, name, price_in]."""
    df = pd.read_excel(filepath, header=None, sheet_name=0)
    layout = detect_layout(df)
    result = parse_cards(df, config) if layout == 'card' else parse_table(df, config)

    if len(result) == 0:
        raise ValueError(
            "Не удалось извлечь ни одной позиции — возможно, структура файла "
            "изменилась. Проверьте, что это прайс ИП Аксёнов."
        )
    return result


# ─── Применение наценок ───────────────────────────────────────────────────────

def apply_pricing(df, group_markups, homeline_prices, default_markup):
    df = df.copy()
    prices_out = []

    for _, row in df.iterrows():
        article = row['article']

        if article in homeline_prices:
            prices_out.append(homeline_prices[article])
            continue

        group = row['group']
        markup_pct = group_markups.get(group)
        if markup_pct is None:
            markup_pct = group_markups.get(group.rstrip())
        if markup_pct is None:
            markup_pct = default_markup

        price = row['price_in'] * (1 + markup_pct / 100)
        prices_out.append(round(price, 0))

    df['price_out'] = prices_out
    return df


# ─── Лист «Прайс клиента» ────────────────────────────────────────────────────

def build_pricelist(df, wb, config, group_order):
    ws = wb.active
    ws.title = 'Прайс клиента'
    company = config['company']

    for letter, w in COL_WIDTHS.items():
        ws.column_dimensions[letter].width = w

    dark_fill = PatternFill('solid', start_color=COLORS['header_bg'])

    # ─── Строка 1: Название компании ───
    ws.merge_cells('A1:F1')
    c = ws['A1']
    c.value = company['name']
    c.font = Font(bold=True, size=22, name='Arial', color=COLORS['header_font'])
    c.fill = dark_fill
    c.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 38.1
    for col in range(2, 7):
        ws.cell(1, col).fill = dark_fill

    # ─── Строка 2: Город ───
    ws.merge_cells('A2:F2')
    c = ws['A2']
    c.value = company['city']
    c.font = Font(size=11, name='Arial', color=COLORS['city_font'])
    c.fill = dark_fill
    c.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[2].height = 21.95
    for col in range(2, 7):
        ws.cell(2, col).fill = dark_fill

    # ─── Строка 3: Телефоны ───
    ws.merge_cells('A3:F3')
    c = ws['A3']
    c.value = company['phones']
    c.font = Font(bold=True, size=10, name='Arial', color=COLORS['phone_font'])
    c.fill = dark_fill
    c.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[3].height = 21.95
    for col in range(2, 7):
        ws.cell(3, col).fill = dark_fill

    # ─── Строка 4: Telegram с гиперссылками ───
    tg_links = company.get('telegram_links', [])

    ws.merge_cells('A4:C4')
    c = ws['A4']
    if tg_links:
        c.value = tg_links[0]['text']
        c.hyperlink = tg_links[0]['url']
    c.font = Font(bold=True, size=12, name='Arial', color=COLORS['tg_font'], underline='single')
    c.fill = dark_fill
    c.alignment = Alignment(horizontal='left', vertical='center')
    for col in range(2, 4):
        ws.cell(4, col).fill = dark_fill

    ws.merge_cells('D4:F4')
    c = ws['D4']
    if len(tg_links) > 1:
        c.value = tg_links[1]['text']
        c.hyperlink = tg_links[1]['url']
    c.font = Font(bold=True, size=12, name='Arial', color=COLORS['tg_font'], underline='single')
    c.fill = dark_fill
    c.alignment = Alignment(horizontal='right', vertical='center')
    for col in range(5, 7):
        ws.cell(4, col).fill = dark_fill
    ws.row_dimensions[4].height = 21.95

    # ─── Строка 5: Заголовки таблицы ───
    headers = ['№', 'Артикул', 'Бренд', 'Наименование', 'Цена (руб.)', 'Заказ (шт.)']
    header_style = {
        'font': Font(bold=True, size=10, name='Arial', color=COLORS['header_font']),
        'fill': PatternFill('solid', start_color=COLORS['header_bg']),
        'alignment': Alignment(horizontal='center', vertical='center', wrap_text=True),
        'border': _border(),
    }
    for ci, h in enumerate(headers, 1):
        _apply(ws.cell(row=5, column=ci, value=h), header_style)

    ws.freeze_panes = 'A6'
    ws.auto_filter.ref = 'A5:F5'

    # ─── Сортировка групп: сначала приоритетные, потом по заданному порядку ───
    max_order = 9999
    unique_groups = df['group'].unique().tolist()
    priority_keys = [p.lower() for p in config.get('priority_groups', [])]

    def group_sort_key(g):
        g_lower = g.lower()
        for i, p in enumerate(priority_keys):
            if p in g_lower:
                return (0, i, g)
        pos = group_order.get(g) or group_order.get(g.rstrip())
        return (1, pos if pos is not None else max_order, g)

    sorted_groups = sorted(unique_groups, key=group_sort_key)

    GEN_GROUP = 'Генераторы (электростанции)'

    # ─── Данные ───
    row = 6
    num = 1

    for group in sorted_groups:
        gdf = df[df['group'] == group].copy()

        if group == GEN_GROUP:
            gdf['_sk'] = gdf['name'].apply(_generator_sort_key)
            gdf = gdf.sort_values('_sk').drop(columns=['_sk'])

        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        gc = ws.cell(row=row, column=1, value=group)
        gc.font = Font(bold=True, size=10, name='Arial', color=COLORS['group_font'])
        gc.fill = PatternFill('solid', start_color=COLORS['group_bg'])
        gc.alignment = Alignment(horizontal='left', vertical='center', indent=1)
        row += 1

        for _, item in gdf.iterrows():
            is_alt = (num % 2 == 0)
            alt_fill = PatternFill('solid', start_color=COLORS['alt_row']) if is_alt else None

            vals = [num, item['article'], item['brand'], item['name'], item['price_out'], '']

            for ci, v in enumerate(vals, 1):
                cell = ws.cell(row=row, column=ci, value=v)
                cell.font = Font(size=10, name='Arial')
                cell.border = _border()
                cell.alignment = Alignment(vertical='center', wrap_text=(ci == 4))
                if ci == 5:
                    cell.alignment = Alignment(horizontal='center', vertical='center')
                    cell.number_format = '#,##0'
                if alt_fill:
                    cell.fill = alt_fill

            row += 1
            num += 1


# ─── Главная функция ──────────────────────────────────────────────────────────

def _save_workbook(wb, output_dir, base_name):
    """Сохраняет книгу в output_dir/base_name.xlsx.
    Если основной файл занят (обычно открыт в Excel) — сохраняет под именем
    с меткой времени, чтобы результат не потерялся и не было аварийного трейсбека.
    Возвращает фактический путь сохранённого файла."""
    primary = Path(output_dir) / f"{base_name}.xlsx"
    try:
        wb.save(primary)
        return str(primary)
    except PermissionError:
        alt = Path(output_dir) / f"{base_name}_{datetime.now():%H-%M-%S}.xlsx"
        try:
            wb.save(alt)
        except PermissionError:
            raise PermissionError(
                f"Не удаётся сохранить в папку output: файл {primary.name} "
                f"открыт в Excel. Закройте его и запустите снова."
            )
        print(f"  Внимание: {primary.name} занят (открыт в Excel) —")
        print(f"  сохранил под именем {alt.name}")
        return str(alt)


def transform(input_path, output_dir=None):
    import unicodedata
    config = load_config()

    print(f"  Читаю файл: {unicodedata.normalize('NFC', Path(input_path).name)}")
    df = read_supplier_price(input_path, config)
    print(f"  Загружено позиций: {len(df)}")

    default_markup = config.get('default_markup', 13)
    print(f"  Наценка: +{default_markup}%")

    df = apply_pricing(df, {}, {}, default_markup)

    wb = openpyxl.Workbook()
    build_pricelist(df, wb, config, {})
    # Лист «Наценки» НЕ добавляется — только 1 лист «Прайс клиента»

    if output_dir is None:
        output_dir = str(BASE_DIR / 'output')
    os.makedirs(output_dir, exist_ok=True)

    prefix = config.get('output_prefix', 'БытТехОпт')
    ts = datetime.today().strftime('%Y%m%d')
    output_path = _save_workbook(wb, output_dir, f"{prefix}_{ts}")
    print(f"  Сохранено: {output_path}")
    return output_path


if __name__ == '__main__':
    if len(sys.argv) < 2:
        print("Использование: python transform.py <прайс_поставщика.xlsx>")
        sys.exit(1)
    transform(sys.argv[1])
