"""
Тесты для Конвертер_Аксенов.
Запуск: python -m pytest test_transform.py -v
"""
import io
import sys
import pytest
import pandas as pd
from pathlib import Path

# Добавляем папку проекта в путь
sys.path.insert(0, str(Path(__file__).parent))

from transform import read_supplier_price, apply_pricing, build_pricelist, extract_brand
import openpyxl


# ─── Вспомогательная функция: создать минимальный .xlsx в памяти ──────────────

def make_xlsx_bytes(rows_data):
    """Табличный формат: 1 строка = 1 товар.
    rows_data — список кортежей (name, price_or_none).
    col 0 = Артикул, col 13 = Номенклатура (N), col 14 = Опт-цена (O)."""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    # Реквизиты + строки-заглушки
    for _ in range(4):
        ws.append([None] * 15)
    # Строка-заголовок (по ней парсер находит колонки)
    hdr = [None] * 15
    hdr[0] = 'Артикул'
    hdr[13] = 'Номенклатура'
    hdr[14] = 'Опт -8% (НАЛ)'
    ws.append(hdr)
    # Данные
    for name, price in rows_data:
        row = [None] * 15
        row[13] = name
        row[14] = price
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _fmt_rub(n):
    """520 -> '520,00 RUB',  1140 -> '1 140,00 RUB' (как в карточном прайсе)."""
    s = f'{int(n):,}'.replace(',', ' ')  # неразрывный пробел-разделитель тысяч
    return f'{s},00 RUB'


def make_card_xlsx_bytes(rows):
    """Карточный формат: 1 товар = блок из нескольких строк.
    rows — список кортежей:
      ('group', 'Название группы')            → строка-заголовок группы (col 1)
      ('item', name, article, price_number)   → карточка товара (3 строки)
    Раскладка карточки повторяет реальный файл: наименование в col 5,
    подписи 'Код'/'Артикул'/'Цена', значения в col 5/9/13, цена текстом '… RUB'."""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    # Реквизиты в col 10 (как в реальном файле)
    for i in range(4):
        rr = [None] * 14
        rr[10] = f'ИП Аксёнов Н.Ю. {i}'
        ws.append(rr)
    ws.append([None] * 14)
    for row in rows:
        if row[0] == 'group':
            rr = [None] * 14
            rr[1] = row[1]
            ws.append(rr)
        else:
            _, name, article, price = row
            rr = [None] * 14; rr[1] = 'НЕТ\nФОТОГРАФИИ'; rr[5] = name; ws.append(rr)
            rr = [None] * 14; rr[5] = 'Код'; rr[9] = 'Артикул'; rr[13] = 'Цена'; ws.append(rr)
            rr = [None] * 14; rr[5] = f'УТ-{article}'; rr[9] = article; rr[13] = _fmt_rub(price); ws.append(rr)
            ws.append([None] * 14)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


CONFIG = {'read_options': {'skiprows': 8}, 'default_markup': 13}


# ─── Тест 1: Артикулы генерируются в формате UT-XXXXXX ────────────────────────

def test_articles_format():
    data = [
        ('Группа товара', None),
        ('Дрель Интерскол 500Вт', 1000),
        ('Шуруповёрт Bosch 18V', 3000),
    ]
    df = read_supplier_price(make_xlsx_bytes(data), CONFIG)
    assert df['article'].tolist() == ['UT-000001', 'UT-000002']


# ─── Тест 2: Артикулы строго последовательны (сквозная нумерация) ─────────────

def test_articles_sequential():
    data = [
        ('Группа 1', None),
        ('Товар A', 100),
        ('Товар B', 200),
        ('Группа 2', None),
        ('Товар C', 300),
    ]
    df = read_supplier_price(make_xlsx_bytes(data), CONFIG)
    assert df['article'].tolist() == ['UT-000001', 'UT-000002', 'UT-000003']


# ─── Тест 3: Группы назначаются правильно ─────────────────────────────────────

def test_group_assignment():
    data = [
        ('Инструменты', None),
        ('Дрель Интерскол', 500),
        ('Генераторы', None),
        ('Генератор ELTI 3кВт', 15000),
        ('Генератор RATO 5кВт', 25000),
    ]
    df = read_supplier_price(make_xlsx_bytes(data), CONFIG)
    assert df.loc[df['article'] == 'UT-000001', 'group'].values[0] == 'Инструменты'
    assert df.loc[df['article'] == 'UT-000002', 'group'].values[0] == 'Генераторы'
    assert df.loc[df['article'] == 'UT-000003', 'group'].values[0] == 'Генераторы'


# ─── Тест 4: Строки без названия пропускаются ─────────────────────────────────

def test_empty_rows_skipped():
    data = [
        (None, None),
        ('', None),
        ('Группа', None),
        (None, 999),       # строка без имени — тоже пропускается
        ('Товар X', 500),
    ]
    df = read_supplier_price(make_xlsx_bytes(data), CONFIG)
    assert len(df) == 1
    assert df.iloc[0]['name'] == 'Товар X'


# ─── Тест 5: Наценка 13% применяется ко всем товарам ─────────────────────────

def test_markup_13_percent():
    df = pd.DataFrame([
        {'article': 'UT-000001', 'group': 'Тест', 'brand': '', 'name': 'Товар', 'price_in': 1000.0},
        {'article': 'UT-000002', 'group': 'Тест', 'brand': '', 'name': 'Товар2', 'price_in': 500.0},
    ])
    result = apply_pricing(df, group_markups={}, homeline_prices={}, default_markup=13)
    assert result.loc[0, 'price_out'] == round(1000 * 1.13, 0)
    assert result.loc[1, 'price_out'] == round(500 * 1.13, 0)


# ─── Тест 6: Наценка не применяется к товарам с фиксированной ценой ──────────

def test_fixed_price_overrides_markup():
    df = pd.DataFrame([
        {'article': 'UT-000001', 'group': 'Тест', 'brand': '', 'name': 'Товар', 'price_in': 1000.0},
    ])
    result = apply_pricing(df, group_markups={}, homeline_prices={'UT-000001': 999.0}, default_markup=13)
    assert result.loc[0, 'price_out'] == 999.0


# ─── Тест 7: Извлечение бренда из названия ───────────────────────────────────

@pytest.mark.parametrize("name, expected", [
    ("Аккумуляторная дрель EINHELL PXC AXXIO",       "EINHELL"),
    ("Аккумулятор ВИТЯЗЬ Тип-М, АКБ-М-20-2.0",       "ВИТЯЗЬ"),
    ("Бетоносмеситель ОПТИМА ББС-130",                "ОПТИМА"),
    ("Генератор CARVER GPG-9500IE 3фазный",           "CARVER"),
    ("Аккумуляторная батарея ELTI только для серии",  "ELTI"),
    ("Шлифмашина угловая ФИОЛЕНТ МШУ1-23-230Б",      "ФИОЛЕНТ"),
    # Прилагательные НЕ должны быть брендом
    ("ЗАРЯДНОЕ устройство 12V ACID",                  "ACID"),
    # Нет бренда → Noname
    ("Гнездо большое сварочное 35-50",                "Noname"),
])
def test_extract_brand(name, expected):
    assert extract_brand(name) == expected, f"name={name!r}"


# ─── Тест 8: Бренды в реальном файле не пустые ───────────────────────────────

# ─── Тест 9: Интеграционный — реальный файл поставщика ───────────────────────

# Ищем xlsx в input/, исключаем временные файлы Excel (~$...)
_input_files = [f for f in (Path(__file__).parent / 'input').glob('*.xlsx')
                if not f.name.startswith('~')]
# Карточный файл помечен 'card' в имени; остальные считаем табличными
REAL_CARD_FILE = next((f for f in _input_files if 'card' in f.name.lower()), None)
REAL_TABLE_FILE = next((f for f in _input_files if 'card' not in f.name.lower()), None)

@pytest.mark.skipif(REAL_TABLE_FILE is None, reason="Табличный файл поставщика не найден в input/")
def test_real_file_parse():
    df = read_supplier_price(str(REAL_TABLE_FILE), CONFIG)  # type: ignore[arg-type]

    assert len(df) > 500, f"Ожидалось >500 позиций, получено {len(df)}"
    assert df['article'].str.match(r'^UT-\d{6}$').all(), "Найдены артикулы неверного формата"
    assert df['name'].str.strip().ne('').all(), "Есть пустые названия товаров"
    assert df['group'].ne('').all(), "Есть товары без группы"
    assert (df['price_in'] > 0).all(), "Есть нулевые или отрицательные цены"
    # Бренды заполнены (не пустые строки)
    assert df['brand'].ne('').all(), "Есть товары с пустым брендом"
    # Должны быть реальные бренды, не только Noname
    assert (df['brand'] != 'Noname').sum() > 100, "Слишком мало распознанных брендов"

    print(f"\n  Позиций: {len(df)}")
    print(f"  Групп:   {df['group'].nunique()}")
    print(f"  Брендов: {df['brand'].nunique()} ({(df['brand'] != 'Noname').sum()} не-Noname)")


# ─── Тест 8: Генераторы идут первой группой в прайсе ─────────────────────────

def test_generators_first_group():
    """Группа, содержащая 'Генераторы', должна идти первой при priority_groups."""
    df = pd.DataFrame([
        {'article': 'UT-000001', 'group': 'Инструменты', 'brand': '', 'name': 'Дрель', 'price_in': 500.0},
        {'article': 'UT-000002', 'group': 'Генераторы (электростанции)', 'brand': '', 'name': 'Генератор 3кВт', 'price_in': 15000.0},
        {'article': 'UT-000003', 'group': 'Бензопилы', 'brand': '', 'name': 'Пила', 'price_in': 8000.0},
    ])
    df = apply_pricing(df, {}, {}, 13)

    cfg = {
        'company': {'name': 'Тест', 'city': '', 'phones': '', 'telegram_links': []},
        'priority_groups': ['Генераторы'],
    }
    wb = openpyxl.Workbook()
    build_pricelist(df, wb, cfg, group_order={})
    ws = wb.active

    # Находим строки-заголовки групп (цвет фона group_bg = BDD7EE)
    group_rows = []
    for row in ws.iter_rows(min_row=6):
        cell = row[0]
        if cell.value and cell.fill and 'BDD7EE' in cell.fill.fgColor.rgb:
            group_rows.append(str(cell.value))

    assert group_rows[0] == 'Генераторы (электростанции)', (
        f"Первая группа должна быть 'Генераторы (электростанции)', получено: '{group_rows[0]}'"
    )


# ─── Тест 10: Парсинг текстовой цены '… RUB' ─────────────────────────────────

@pytest.mark.parametrize("text, expected", [
    ('520,00 RUB', 520.0),
    ('1 140,00 RUB', 1140.0),
    ('3 995,00 RUB', 3995.0),
    ('160,00 RUB', 160.0),
    ('12 500,50 RUB', 12500.5),
])
def test_parse_price_text(text, expected):
    from transform import _parse_price_text
    assert _parse_price_text(text) == expected


# ─── Тест 11: Определение макета файла (табличный / карточный) ────────────────

def test_detect_layout_table():
    from transform import detect_layout
    df = pd.read_excel(make_xlsx_bytes([('Группа', None), ('Товар', 500)]), header=None)
    assert detect_layout(df) == 'table'


def test_detect_layout_card():
    from transform import detect_layout
    rows = [('group', 'Группа'), ('item', 'Товар', 'A1', 520)]
    df = pd.read_excel(make_card_xlsx_bytes(rows), header=None)
    assert detect_layout(df) == 'card'


# ─── Тест 12: Разбор карточного формата ──────────────────────────────────────

def test_card_parse_basic():
    rows = [
        ('group', 'Аккумуляторный инструмент'),
        ('item', 'Аккумулятор ВИТЯЗЬ Тип-М 20В', '18037001', 520),
        ('item', 'Дрель-шуруповерт ВИТЯЗЬ ДА-201', '18012031', 3305),
        ('group', 'Генераторы (электростанции)'),
        ('item', 'Генератор CARVER GPG 3кВт', 'ГЕН-1', 15000),
    ]
    df = read_supplier_price(make_card_xlsx_bytes(rows), CONFIG)

    assert len(df) == 3
    assert df.iloc[0]['name'] == 'Аккумулятор ВИТЯЗЬ Тип-М 20В'
    assert df.iloc[0]['price_in'] == 520.0
    assert df.iloc[0]['group'] == 'Аккумуляторный инструмент'
    assert df.iloc[2]['group'] == 'Генераторы (электростанции)'
    # Артикулы обезличиваются (UT-…), а не берутся из файла
    assert df['article'].tolist() == ['UT-000001', 'UT-000002', 'UT-000003']
    # Бренд извлекается из наименования
    assert df.iloc[0]['brand'] == 'ВИТЯЗЬ'


def test_card_thousands_price():
    """Цена с разделителем тысяч '1 140,00 RUB' парсится в 1140.0."""
    rows = [('group', 'Г'), ('item', 'Товар X', 'A1', 1140)]
    df = read_supplier_price(make_card_xlsx_bytes(rows), CONFIG)
    assert df.iloc[0]['price_in'] == 1140.0


# ─── Тест 13: Табличный парсер находит колонки по заголовку (не по индексу) ───

def test_table_autodetect_shifted_columns():
    """Колонки определяются по строке-заголовку, даже если сдвинуты."""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append([None] * 5)
    ws.append([None] * 5)
    ws.append([None, 'Наименование', 'Цена', None, None])   # заголовок: колонки B, C
    ws.append([None, 'Дрель Интерскол', 1000, None, None])
    ws.append([None, 'Генераторы', None, None, None])        # группа (нет цены)
    ws.append([None, 'Генератор CARVER 3кВт', 15000, None, None])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    df = read_supplier_price(buf, {'read_options': {}, 'default_markup': 13})
    assert len(df) == 2
    assert df.iloc[0]['name'] == 'Дрель Интерскол'
    assert df.iloc[0]['price_in'] == 1000
    assert df.iloc[1]['group'] == 'Генераторы'


# ─── Тест 14: Нераспознанная структура даёт понятную ошибку ───────────────────

def test_unrecognized_structure_raises():
    """Файл без цен и заголовков → явная ошибка, а не тихие 0 позиций."""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    for _ in range(6):
        ws.append(['текст', 'ещё текст', 'без цифр'])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    with pytest.raises(ValueError):
        read_supplier_price(buf, CONFIG)


# ─── Тест 15: Интеграционный — реальный карточный файл ───────────────────────

@pytest.mark.skipif(REAL_CARD_FILE is None, reason="Карточный файл не найден в input/")
def test_real_card_file_parse():
    df = read_supplier_price(str(REAL_CARD_FILE), CONFIG)  # type: ignore[arg-type]

    assert len(df) > 500, f"Ожидалось >500 позиций, получено {len(df)}"
    assert df['article'].str.match(r'^UT-\d{6}$').all(), "Найдены артикулы неверного формата"
    assert df['name'].str.strip().ne('').all(), "Есть пустые названия товаров"
    assert df['group'].ne('').all(), "Есть товары без группы"
    assert (df['price_in'] > 0).all(), "Есть нулевые или отрицательные цены"
    assert df['brand'].ne('').all(), "Есть товары с пустым брендом"

    print(f"\n  [карточный] Позиций: {len(df)}")
    print(f"  [карточный] Групп:   {df['group'].nunique()}")


# ─── Тест 16: detect_layout на реальных файлах ───────────────────────────────

@pytest.mark.skipif(REAL_TABLE_FILE is None, reason="Табличный файл не найден")
def test_real_table_layout_detected():
    from transform import detect_layout
    df = pd.read_excel(str(REAL_TABLE_FILE), header=None)
    assert detect_layout(df) == 'table'


@pytest.mark.skipif(REAL_CARD_FILE is None, reason="Карточный файл не найден")
def test_real_card_layout_detected():
    from transform import detect_layout
    df = pd.read_excel(str(REAL_CARD_FILE), header=None)
    assert detect_layout(df) == 'card'
